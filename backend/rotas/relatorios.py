from datetime import date, timedelta
from dateutil.relativedelta import relativedelta
from io import BytesIO
from flask import Blueprint, render_template, send_file, make_response
from flask_login import login_required
from sqlalchemy import func
from models import db, Aluno, Matricula, Mensalidade, Pagamento

relatorios_bp = Blueprint("relatorios", __name__)

MESES = [
    "", "Janeiro", "Fevereiro", "Marco", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
]


def _calcularPerfilPagamento(aluno):
    pagamentos = (
        db.session.query(Pagamento)
        .join(Mensalidade)
        .join(Matricula)
        .filter(Matricula.aluno_id == aluno.id, Pagamento.data_pagamento.isnot(None))
        .all()
    )
    if not pagamentos:
        return {"dias_medio": 0, "tendencia": "Sem dados", "status": "sem_dados", "total_pagamentos": 0}

    dias_atraso = []
    for p in pagamentos:
        if p.mensalidade and p.data_pagamento and p.mensalidade.data_vencimento:
            delta = (p.data_pagamento - p.mensalidade.data_vencimento).days
            dias_atraso.append(delta)

    if not dias_atraso:
        return {"dias_medio": 0, "tendencia": "Sem dados", "status": "sem_dados", "total_pagamentos": len(pagamentos)}

    media = sum(dias_atraso) / len(dias_atraso)

    if media <= 0:
        status = "pontual"
        tendencia = "✅ Pontual"
    elif media <= 3:
        status = "levemente_atrasado"
        tendencia = "⚠️ Levemente atrasado"
    else:
        status = "sempre_atrasado"
        tendencia = "🔴 Sempre atrasado"

    return {
        "dias_medio": round(media, 1),
        "tendencia": tendencia,
        "status": status,
        "total_pagamentos": len(pagamentos),
    }


def _obterDadosRelatorios():
    hoje = date.today()
    mes_atual = hoje.month
    ano_atual = hoje.year
    proximo_mes = hoje + relativedelta(months=1)

    # 1. FATURAMENTO DO MES ATUAL
    pagamentos_mes = (
        db.session.query(Pagamento)
        .join(Mensalidade)
        .filter(
            func.extract("month", Pagamento.data_pagamento) == mes_atual,
            func.extract("year", Pagamento.data_pagamento) == ano_atual,
        )
        .all()
    )
    total_faturamento = sum(float(p.valor_pago) for p in pagamentos_mes)

    # 2. INADIMPLENCIA
    mensalidades_atrasadas = (
        Mensalidade.query.filter_by(paga=False)
        .filter(Mensalidade.data_vencimento < hoje)
        .order_by(Mensalidade.data_vencimento)
        .all()
    )
    total_inadimplencia = sum(float(m.valor) for m in mensalidades_atrasadas)

    # 3. ESTIMATIVA PROXIMO MES - baseada em matriculas ativas
    matriculas_ativas = Matricula.query.filter_by(ativa=True).all()
    estimativas = []
    for mat in matriculas_ativas:
        # Verificar se ja tem mensalidade para proximo mes
        mensalidade_proximo = None
        for m in mat.mensalidades:
            if m.data_vencimento.month == proximo_mes.month and m.data_vencimento.year == proximo_mes.year:
                mensalidade_proximo = m
                break

        if mensalidade_proximo:
            if not mensalidade_proximo.paga:
                estimativas.append({
                    "aluno": mat.aluno.nome,
                    "plano": mat.plano.nome,
                    "valor": float(mensalidade_proximo.valor),
                    "vencimento": mensalidade_proximo.data_vencimento,
                    "status": "pendente",
                })
        else:
            # Mensalidade ainda nao gerada, mas matricula ativa = estimar
            estimativas.append({
                "aluno": mat.aluno.nome,
                "plano": mat.plano.nome,
                "valor": float(mat.plano.valor / mat.plano.duracao_meses),
                "vencimento": hoje + relativedelta(months=1, day=mat.data_inicio.day),
                "status": "estimado",
            })

    total_estimativa = sum(e["valor"] for e in estimativas)

    # 4. ALUNOS COM AVISO (faltam 7 dias ou menos)
    data_limite = hoje + timedelta(days=7)
    mensalidades_aviso = (
        Mensalidade.query.filter_by(paga=False)
        .filter(
            Mensalidade.data_vencimento >= hoje,
            Mensalidade.data_vencimento <= data_limite,
        )
        .order_by(Mensalidade.data_vencimento)
        .all()
    )

    # 5. PERFIL DE PAGAMENTO DOS ALUNOS
    alunos_ativos = Aluno.query.filter_by(ativo=True).all()
    perfis = []
    for aluno in alunos_ativos:
        perfil = _calcularPerfilPagamento(aluno)
        if perfil["total_pagamentos"] > 0:
            perfis.append({"aluno": aluno.nome, **perfis_append(perfil)})
    perfis.sort(key=lambda x: x["dias_medio"], reverse=True)

    # 6. DADOS PARA GRAFICOS
    # Ultimos 6 meses de faturamento
    faturamento_mensal = []
    for i in range(5, -1, -1):
        data_ref = hoje - relativedelta(months=i)
        pagamentos = (
            db.session.query(Pagamento)
            .join(Mensalidade)
            .filter(
                func.extract("month", Pagamento.data_pagamento) == data_ref.month,
                func.extract("year", Pagamento.data_pagamento) == data_ref.year,
            )
            .all()
        )
        total = sum(float(p.valor_pago) for p in pagamentos)
        faturamento_mensal.append({"mes": MESES[data_ref.month], "valor": total})

    # Status das mensalidades do mes atual
    todas_mensalidades_mes = Mensalidade.query.filter(
        func.extract("month", Mensalidade.data_vencimento) == mes_atual,
        func.extract("year", Mensalidade.data_vencimento) == ano_atual,
    ).all()
    status_count = {"pagas": 0, "pendentes": 0, "atrasadas": 0}
    for m in todas_mensalidades_mes:
        if m.paga:
            status_count["pagas"] += 1
        elif m.esta_atrasada:
            status_count["atrasadas"] += 1
        else:
            status_count["pendentes"] += 1

    return {
        "pagamentos_mes": pagamentos_mes,
        "total_faturamento": total_faturamento,
        "mensalidades_atrasadas": mensalidades_atrasadas,
        "total_inadimplencia": total_inadimplencia,
        "estimativas": estimativas,
        "total_estimativa": total_estimativa,
        "mensalidades_aviso": mensalidades_aviso,
        "perfis": perfis,
        "faturamento_mensal": faturamento_mensal,
        "status_count": status_count,
        "mes_atual": mes_atual,
        "ano_atual": ano_atual,
        "proximo_mes": proximo_mes,
        "hoje": hoje,
    }


def perfis_append(perfil):
    return {
        "dias_medio": perfil["dias_medio"],
        "tendencia": perfil["tendencia"],
        "status": perfil["status"],
        "total_pagamentos": perfil["total_pagamentos"],
    }


@relatorios_bp.route("/")
@login_required
def index():
    dados = _obterDadosRelatorios()
    return render_template(
        "relatorios/index.html",
        meses=MESES,
        **dados,
    )


@relatorios_bp.route("/pdf")
@login_required
def pdf():
    dados = _obterDadosRelatorios()
    html = render_template(
        "relatorios/pdf.html",
        meses=MESES,
        **dados,
    )
    from xhtml2pdf import pisa
    result = BytesIO()
    pdf = pisa.pisaDocument(html, result)
    if not pdf.err:
        response = make_response(result.getvalue())
        response.headers["Content-Type"] = "application/pdf"
        response.headers["Content-Disposition"] = "attachment; filename=relatorio_academia_vip.pdf"
        return response
    return "Erro ao gerar PDF", 500


@relatorios_bp.route("/excel")
@login_required
def excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    dados = _obterDadosRelatorios()
    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
    gold_font = Font(bold=True, color="D4AF37")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    def estilizar_cabecalho(ws, cols):
        for col in range(1, cols + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

    # ABA 1: Faturamento
    ws1 = wb.active
    ws1.title = "Faturamento"
    ws1.append(["Data Pagamento", "Aluno", "Plano", "Valor", "Forma"])
    estilizar_cabecalho(ws1, 5)
    for p in dados["pagamentos_mes"]:
        ws1.append([
            p.data_pagamento.strftime("%d/%m/%Y"),
            p.mensalidade.matricula.aluno.nome,
            p.mensalidade.matricula.plano.nome,
            float(p.valor_pago),
            p.forma_pagamento.replace("_", " ").title(),
        ])
    ws1.append([])
    ws1.append(["TOTAL", "", "", dados["total_faturamento"], ""])
    ws1.cell(row=ws1.max_row, column=1).font = gold_font
    ws1.cell(row=ws1.max_row, column=4).font = gold_font
    for col in ws1.columns:
        ws1.column_dimensions[col[0].column_letter].width = 20

    # ABA 2: Inadimplencia
    ws2 = wb.create_sheet("Inadimplencia")
    ws2.append(["Aluno", "Plano", "Valor", "Vencimento", "Dias em Atraso"])
    estilizar_cabecalho(ws2, 5)
    for m in dados["mensalidades_atrasadas"]:
        dias = (dados["hoje"] - m.data_vencimento).days
        ws2.append([
            m.matricula.aluno.nome,
            m.matricula.plano.nome,
            float(m.valor),
            m.data_vencimento.strftime("%d/%m/%Y"),
            dias,
        ])
    ws2.append([])
    ws2.append(["TOTAL", "", dados["total_inadimplencia"], "", ""])
    ws2.cell(row=ws2.max_row, column=1).font = gold_font
    ws2.cell(row=ws2.max_row, column=3).font = gold_font
    for col in ws2.columns:
        ws2.column_dimensions[col[0].column_letter].width = 20

    # ABA 3: Estimativa Proximo Mes
    ws3 = wb.create_sheet(f"Estimativa {MESES[dados['proximo_mes'].month]}")
    ws3.append(["Aluno", "Plano", "Valor", "Vencimento", "Status"])
    estilizar_cabecalho(ws3, 5)
    for e in dados["estimativas"]:
        ws3.append([
            e["aluno"],
            e["plano"],
            e["valor"],
            e["vencimento"].strftime("%d/%m/%Y"),
            "Estimado" if e["status"] == "estimado" else "Pendente",
        ])
    ws3.append([])
    ws3.append(["TOTAL", "", dados["total_estimativa"], "", ""])
    ws3.cell(row=ws3.max_row, column=1).font = gold_font
    ws3.cell(row=ws3.max_row, column=3).font = gold_font
    for col in ws3.columns:
        ws3.column_dimensions[col[0].column_letter].width = 20

    # ABA 4: Perfil de Pagamento
    ws4 = wb.create_sheet("Perfil de Pagamento")
    ws4.append(["Aluno", "Media Dias Atraso", "Tendencia", "Total Pagamentos"])
    estilizar_cabecalho(ws4, 4)
    for p in dados["perfis"]:
        ws4.append([
            p["aluno"],
            p["dias_medio"],
            p["tendencia"],
            p["total_pagamentos"],
        ])
    for col in ws4.columns:
        ws4.column_dimensions[col[0].column_letter].width = 25

    result = BytesIO()
    wb.save(result)
    result.seek(0)
    response = make_response(result.getvalue())
    response.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    response.headers["Content-Disposition"] = "attachment; filename=relatorio_academia_vip.xlsx"
    return response
